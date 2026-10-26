"""Small, testable primitives for the Sprint 2 document-processing vertical slice."""

from __future__ import annotations

import hashlib
import re
import subprocess
import tempfile
import unicodedata
from collections import deque
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Iterator
from zipfile import ZipFile

from docx import Document as DocxDocument
from pypdf import PdfReader

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


class PipelineError(Exception):
    def __init__(self, code: str, message: str, stage: str, status_code: int = 422):
        super().__init__(message)
        self.code = code
        self.stage = stage
        self.status_code = status_code


@dataclass(frozen=True)
class TextUnit:
    text: str
    page_number: int | None = None
    source_locator: str | None = None


@dataclass
class ExtractedDocument:
    text: str
    page_count: int | None
    metadata: dict[str, Any]


def _normalized_heading(value: str) -> str:
    value = unicodedata.normalize('NFKC', value)
    value = re.sub(r'\.{2,}\s*\d+\s*$', '', value)
    return re.sub(r'\s+', ' ', value).strip(' .\t')


def _section_key(order_index: int) -> str:
    return f'section-{order_index:04d}'


def _complete_section_paths(sections: list[dict[str, Any]], final_page: int | None = None) -> list[dict[str, Any]]:
    """Resolve parent paths and source ranges without asking the LLM to invent document structure."""
    by_key = {section['section_key']: section for section in sections}
    for section in sections:
        path: list[str] = []
        cursor: dict[str, Any] | None = section
        while cursor:
            path.insert(0, cursor['title'])
            cursor = by_key.get(cursor.get('parent_key'))
        section['heading_path'] = path

    for index, section in enumerate(sections):
        if section.get('page_start') is None:
            section['page_end'] = None
            continue
        next_boundary = next((
            candidate.get('page_start')
            for candidate in sections[index + 1:]
            if candidate.get('page_start') is not None and candidate['level'] <= section['level']
        ), None)
        section['page_end'] = max(section['page_start'], next_boundary - 1) if next_boundary else final_page
    return sections


def _pdf_native_outline(reader: PdfReader, page_limit: int) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []

    def visit(entries: list[Any], level: int = 1, parent_key: str | None = None) -> None:
        last_key = parent_key
        for entry in entries:
            if isinstance(entry, list):
                visit(entry, min(5, level + 1), last_key)
                continue
            title = _normalized_heading(str(getattr(entry, 'title', '') or ''))
            if not title:
                continue
            try:
                page_start = reader.get_destination_page_number(entry) + 1
            except Exception:
                page_start = None
            if page_start is not None and page_start > page_limit:
                continue
            key = _section_key(len(sections) + 1)
            sections.append({
                'section_key': key,
                'parent_key': parent_key,
                'level': level,
                'order_index': len(sections) + 1,
                'title': title,
                'source_locator': f'page {page_start}' if page_start else None,
                'page_start': page_start,
                'detection_source': 'pdf_bookmark',
                'confidence': 0.98,
            })
            last_key = key

    try:
        visit(list(reader.outline or []))
    except Exception:
        return []
    return _complete_section_paths(sections, page_limit)


def _docx_native_outline(document: DocxDocument) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    parents: dict[int, str] = {}
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        title = _normalized_heading(paragraph.text)
        style_name = str(getattr(paragraph.style, 'name', '') or '')
        match = re.search(r'(?:heading|tiêu đề)\s*(\d+)', style_name, re.IGNORECASE)
        if not title or not match:
            continue
        level = min(5, max(1, int(match.group(1))))
        parent_key = parents.get(level - 1)
        key = _section_key(len(sections) + 1)
        sections.append({
            'section_key': key,
            'parent_key': parent_key,
            'level': level,
            'order_index': len(sections) + 1,
            'title': title,
            'source_locator': f'paragraph {paragraph_index}',
            'page_start': None,
            'detection_source': 'docx_style',
            'confidence': 0.98,
        })
        parents[level] = key
        parents = {key_level: value for key_level, value in parents.items() if key_level <= level}
    return _complete_section_paths(sections)


def extract_document(content: bytes, extension: str, max_pages: int | None = None) -> ExtractedDocument:
    extension = extension.lower().lstrip('.')
    if extension == 'pdf':
        return _extract_pdf(content, max_pages=max_pages)
    if extension == 'docx':
        return _extract_docx(content)
    if extension == 'doc':
        return _extract_doc(content)
    raise PipelineError('UNSUPPORTED_FILE_TYPE', 'Only PDF, DOCX and DOC are supported.', 'extracting')


def _extract_pdf(content: bytes, max_pages: int | None = None) -> ExtractedDocument:
    if max_pages is not None and max_pages <= 0:
        raise ValueError('max_pages must be positive or None')
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise PipelineError('PDF_ENCRYPTED', 'Password-protected PDFs cannot be processed.', 'extracting')
        pages = []
        total_pages = len(reader.pages)
        limit = total_pages if (max_pages is None or max_pages <= 0) else min(total_pages, max_pages)
        for index in range(limit):
            page = reader.pages[index]
            page_text = (page.extract_text() or '').strip()
            if page_text:
                pages.append(f'[Page {index + 1}]\n{page_text}')
    except PipelineError:
        raise
    except Exception as exception:
        raise PipelineError('PDF_EXTRACTION_FAILED', 'The PDF could not be read as text.', 'extracting') from exception

    text = clean_text('\n\n'.join(pages))
    if not text:
        raise PipelineError('NO_EXTRACTABLE_TEXT', 'No extractable text was found. Scanned PDFs need OCR, which is not included in Sprint 2.', 'extracting')
    sections = _pdf_native_outline(reader, limit)
    structure_source = 'pdf_bookmark'
    if len(sections) < 2:
        sections = infer_document_outline(text, final_page=limit)
        structure_source = 'inferred_heading' if sections else 'unstructured'
    return ExtractedDocument(
        text=text,
        page_count=limit,
        metadata={
            'parser': 'pypdf', 'extracted_pages': limit, 'total_pages': total_pages,
            'structure_source': structure_source, 'outline': sections,
        }
    )


def _extract_docx(content: bytes) -> ExtractedDocument:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise PipelineError('DOCX_TOO_LARGE', 'Expanded DOCX exceeds the safe limit.', 'extracting')
        document = DocxDocument(BytesIO(content))
    except PipelineError:
        raise
    except Exception as exception:
        raise PipelineError('DOCX_EXTRACTION_FAILED', 'The DOCX file could not be read.', 'extracting') from exception

    parts = [f'[Paragraph {index}]\n{paragraph.text}' for index, paragraph in enumerate(document.paragraphs, start=1) if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        rows = [' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if any(rows):
            parts.append(f'[Table {table_index}]\n' + '\n'.join(rows))
    text = clean_text('\n\n'.join(parts))
    if not text:
        raise PipelineError('NO_EXTRACTABLE_TEXT', 'No extractable text was found in this DOCX document.', 'extracting')
    sections = _docx_native_outline(document)
    structure_source = 'docx_style'
    if len(sections) < 2:
        sections = infer_document_outline(text)
        structure_source = 'inferred_heading' if sections else 'unstructured'
    return ExtractedDocument(
        text=text,
        page_count=None,
        metadata={
            'parser': 'python-docx', 'table_count': len(document.tables),
            'structure_source': structure_source, 'outline': sections,
        },
    )


def _extract_doc(content: bytes) -> ExtractedDocument:
    try:
        with tempfile.NamedTemporaryFile(suffix='.doc', delete=False) as source:
            source.write(content)
            source_path = source.name
        try:
            result = subprocess.run(['antiword', '-m', 'UTF-8.txt', source_path], capture_output=True, text=True, timeout=30, check=False)
        finally:
            Path(source_path).unlink(missing_ok=True)
    except FileNotFoundError as exception:
        raise PipelineError('DOC_PARSER_UNAVAILABLE', 'The legacy DOC parser is not installed in the AI service.', 'extracting', 503) from exception
    except subprocess.TimeoutExpired as exception:
        raise PipelineError('DOC_EXTRACTION_TIMEOUT', 'Legacy DOC extraction exceeded the 30-second limit.', 'extracting', 422) from exception
    except Exception as exception:
        raise PipelineError('DOC_EXTRACTION_FAILED', 'The DOC file could not be read.', 'extracting') from exception

    if result.returncode != 0:
        raise PipelineError('DOC_EXTRACTION_FAILED', 'The legacy DOC file could not be read by the parser.', 'extracting')
    text = clean_text(result.stdout)
    if not text:
        raise PipelineError('NO_EXTRACTABLE_TEXT', 'No extractable text was found in this DOC document.', 'extracting')
    sections = infer_document_outline(text)
    return ExtractedDocument(text=text, page_count=None, metadata={
        'parser': 'antiword',
        'structure_source': 'inferred_heading' if sections else 'unstructured',
        'outline': sections,
    })


def clean_text(value: str) -> str:
    value = unicodedata.normalize('NFKC', value)
    value = value.replace('\r\n', '\n').replace('\r', '\n')
    value = ''.join(
        char for char in value
        if char in '\n\t' or not unicodedata.category(char).startswith('C')
    ).replace('\ufffd', '')
    value = re.sub(
        r'(?i)(?<![\w.+-])[\w.!#$%&\x27*+/=?^`{|}~-]+@[\w-]+(?:\.[\w-]+)+',
        '[redacted-email]', value,
    )
    value = re.sub(r'(?<=[A-Za-z0-9\)\}\| ])\s*È\s*(?=[A-Za-z0-9\(\{\| ])', ' ∪ ', value)
    value = re.sub(r'(?<=[A-Za-z0-9\)\}\| ])\s*Ç\s*(?=[A-Za-z0-9\(\{\| ])', ' ∩ ', value)
    value = re.sub(r'(?<=[A-Za-z0-9\)\}\| ])\s*Î\s*(?=[A-Za-z0-9\(\{\| ])', ' ∈ ', value)
    value = re.sub(r'(?<=[A-Za-z0-9\)\}\| ])\s*Ï\s*(?=[A-Za-z0-9\(\{\| ])', ' ∉ ', value)
    value = re.sub(r'(?<=[A-Za-z0-9\)\}\| ])\s*Ì\s*(?=[A-Za-z0-9\(\{\| ])', ' ⊂ ', value)
    value = re.sub(r'(?<=[A-Za-z0-9\)\}\| ])\s*Í\s*(?=[A-Za-z0-9\(\{\| ])', ' ⊆ ', value)
    value = re.sub(r'[^\S\n]+', ' ', value)
    value = re.sub(r' *\n *', '\n', value)
    return re.sub(r'\n{3,}', '\n\n', value).strip()


SOURCE_MARKER = re.compile(r'^\[(Page|Paragraph|Table|Doc Segment) (\d+)(?: Row \d+)?\]\s*', re.IGNORECASE)
NUMBERED_HEADING = re.compile(r'^(?:\d+(?:\.\d+){0,4}|[IVXLCDM]+)[.)]?\s+\S+', re.IGNORECASE)
NAMED_HEADING = re.compile(r'^(?:chương|chapter|phần|part|mục|section|bài)\s+[\wIVXLCDM.-]+', re.IGNORECASE)


def _heading_level(value: str) -> int:
    normalized = value.strip()
    if re.match(r'^(?:chương|chapter|phần|part)\b', normalized, re.IGNORECASE):
        return 1
    numbered = re.match(r'^(\d+(?:\.\d+)*)', normalized)
    if numbered:
        return min(5, max(2, numbered.group(1).count('.') + 1))
    return 2


def _looks_like_heading(value: str) -> bool:
    candidate = re.sub(r'\s+', ' ', value).strip()
    words = candidate.split()
    if not candidate or len(candidate) > 160 or len(words) > 18:
        return False
    if NAMED_HEADING.match(candidate) or NUMBERED_HEADING.match(candidate):
        return True
    letters = [character for character in candidate if character.isalpha()]
    uppercase_ratio = (sum(character.isupper() for character in letters) / len(letters)) if letters else 0
    return len(words) >= 2 and uppercase_ratio >= 0.75 and not re.search(r'[.!?…:]$', candidate)


def infer_document_outline(text: str, final_page: int | None = None) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    by_title: dict[str, int] = {}
    parents: dict[int, str] = {}
    for block in [part.strip() for part in re.split(r'\n{2,}', text) if part.strip()]:
        marker = SOURCE_MARKER.match(block)
        locator = f'{marker.group(1).lower()} {marker.group(2)}' if marker else None
        page_start = _page_number(locator)
        body = block[marker.end():].strip() if marker else block
        for raw_line in body.split('\n'):
            title = _normalized_heading(raw_line)
            if not _looks_like_heading(title):
                continue
            normalized = re.sub(r'\W+', ' ', title, flags=re.UNICODE).strip().casefold()
            if normalized in by_title:
                existing = sections[by_title[normalized]]
                existing['source_locator'] = locator or existing.get('source_locator')
                existing['page_start'] = page_start if page_start is not None else existing.get('page_start')
                continue
            level = _heading_level(title)
            parent_key = parents.get(level - 1)
            key = _section_key(len(sections) + 1)
            sections.append({
                'section_key': key,
                'parent_key': parent_key,
                'level': level,
                'order_index': len(sections) + 1,
                'title': title,
                'source_locator': locator,
                'page_start': page_start,
                'detection_source': 'inferred_heading',
                'confidence': 0.72 if NAMED_HEADING.match(title) or NUMBERED_HEADING.match(title) else 0.58,
            })
            by_title[normalized] = len(sections) - 1
            parents[level] = key
            parents = {key_level: value for key_level, value in parents.items() if key_level <= level}
    return _complete_section_paths(sections, final_page)


def _sentences(value: str) -> list[str]:
    normalized = re.sub(r'\s+', ' ', value).strip()
    if not normalized:
        return []
    parts = [part.strip() for part in re.split(r'(?<=[.!?…])\s+(?=[A-ZÀ-Ỹ0-9])', normalized) if part.strip()]
    return parts or [normalized]


def _page_number(locator: str | None) -> int | None:
    if not locator:
        return None
    match = re.match(r'page\s+(\d+)', locator, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _locator_number(locator: str | None) -> tuple[str | None, int | None]:
    if not locator:
        return None, None
    match = re.match(r'(page|paragraph|table|doc segment)\s+(\d+)', locator, re.IGNORECASE)
    return (match.group(1).lower(), int(match.group(2))) if match else (None, None)


def attach_document_sections(chunks: list[dict[str, Any]], sections: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not sections:
        return chunks
    by_title = {
        re.sub(r'\W+', ' ', section['title'], flags=re.UNICODE).strip().casefold(): section
        for section in sections
    }
    for chunk in chunks:
        selected = None
        for title in reversed(chunk.get('heading_path') or []):
            normalized = re.sub(r'\W+', ' ', title, flags=re.UNICODE).strip().casefold()
            if normalized in by_title:
                selected = by_title[normalized]
                break
        if selected is None:
            locator_type, locator_number = _locator_number(chunk.get('source_start') or chunk.get('source_locator'))
            candidates = []
            for section in sections:
                section_type, section_number = _locator_number(section.get('source_locator'))
                if locator_number is None or section_number is None or locator_type != section_type:
                    continue
                if section_number <= locator_number:
                    end = section.get('page_end') if locator_type == 'page' else None
                    if end is None or locator_number <= end:
                        candidates.append(section)
            if candidates:
                selected = max(candidates, key=lambda item: (item['level'], item['order_index']))
        if selected:
            chunk['section_key'] = selected['section_key']
            chunk['section_title'] = selected['title']
            chunk['heading_path'] = selected.get('heading_path') or [selected['title']]
            chunk['chunking_strategy'] = 'sentence-outline-v3'
        else:
            chunk['section_key'] = None
    return chunks


def chunk_text(text: str, target_words: int = 420, overlap_words: int = 60) -> list[dict[str, Any]]:
    if target_words <= 0 or not 0 <= overlap_words < target_words:
        raise ValueError('target_words must be > 0 and 0 <= overlap_words < target_words')
    blocks = [part.strip() for part in re.split(r'\n{2,}', text) if part.strip()]
    units: list[dict[str, Any]] = []
    heading_path: list[str] = []

    for block in blocks:
        marker = SOURCE_MARKER.match(block)
        locator = f'{marker.group(1).lower()} {marker.group(2)}' if marker else None
        body = block[marker.end():].strip() if marker else block
        for line in [line.strip() for line in body.split('\n') if line.strip()]:
            if _looks_like_heading(line):
                level = _heading_level(line)
                heading_path = heading_path[:max(0, level - 1)]
                heading_path.append(re.sub(r'\s+', ' ', line).strip())
                continue
            for sentence in _sentences(line):
                units.append({'text': sentence, 'locator': locator, 'heading_path': list(heading_path)})

    chunks: list[dict[str, Any]] = []
    current: list[dict[str, Any]] = []
    current_words = 0

    def append_current() -> None:
        nonlocal current, current_words
        if not current:
            return
        active_path = next((unit['heading_path'] for unit in reversed(current) if unit['heading_path']), [])
        body = ' '.join(unit['text'] for unit in current).strip()
        heading_prefix = f"Section: {' > '.join(active_path)}\n\n" if active_path else ''
        content = f'{heading_prefix}{body}'.strip()
        locators = list(dict.fromkeys(unit['locator'] for unit in current if unit['locator']))
        first_locator = locators[0] if locators else None
        last_locator = locators[-1] if locators else None
        first_page, last_page = _page_number(first_locator), _page_number(last_locator)
        source_locator = first_locator
        if first_page is not None and last_page is not None and first_page != last_page:
            source_locator = f'page {first_page}-{last_page}'
        chunks.append({
            'chunk_index': len(chunks),
            'content': content,
            'source_locator': source_locator,
            'source_start': first_locator,
            'source_end': last_locator,
            'page_start': first_page,
            'page_end': last_page,
            'section_title': active_path[-1] if active_path else None,
            'heading_path': active_path,
            'chunking_strategy': 'sentence-structure-v2',
            'token_estimate': max(1, round(len(content.split()) * 1.3)),
            'content_hash': hashlib.sha256(content.encode('utf-8')).hexdigest(),
        })

        if overlap_words <= 0:
            current, current_words = [], 0
            return
        overlap: list[dict[str, Any]] = []
        overlap_count = 0
        for unit in reversed(current):
            overlap.insert(0, unit)
            overlap_count += len(unit['text'].split())
            if overlap_count >= overlap_words:
                break
        current, current_words = overlap, overlap_count

    for unit in units:
        words = unit['text'].split()
        if current and current_words + len(words) > target_words:
            append_current()
        if len(words) > target_words:
            step = max(1, target_words - overlap_words)
            for offset in range(0, len(words), step):
                segment = words[offset:offset + target_words]
                if not segment:
                    break
                current = [{**unit, 'text': ' '.join(segment)}]
                current_words = len(segment)
                append_current()
                if offset + target_words >= len(words):
                    break
            continue
        current.append(unit)
        current_words += len(words)
    append_current()
    return chunks


def iter_chunks(units: Iterable[TextUnit], target_words: int = 500,
                overlap_words: int = 50) -> Iterator[dict[str, Any]]:
    """Streaming chunker for TextUnit sequences."""
    if target_words <= 0 or not 0 <= overlap_words < target_words:
        raise ValueError('Require target_words > 0 and 0 <= overlap_words < target_words')
    window: deque[tuple[str, int | None, str | None]] = deque()
    chunk_index = 0
    new_words = 0

    def build_chunk() -> dict[str, Any]:
        content = ' '.join(item[0] for item in window)
        pages = list(dict.fromkeys(item[1] for item in window if item[1] is not None))
        locators = list(dict.fromkeys(item[2] for item in window if item[2]))
        return {
            'chunk_index': chunk_index, 'content': content,
            'word_count': len(window), 'token_estimate': max(1, round(len(window) * 1.3)),
            'content_hash': hashlib.sha256(content.encode('utf-8')).hexdigest(),
            'page_number': pages[0] if pages else None,
            'page_numbers': pages, 'source_locator': locators[0] if locators else None,
            'source_locators': locators,
        }

    for unit in units:
        for match in re.finditer(r'\S+', clean_text(unit.text)):
            window.append((match.group(), unit.page_number, unit.source_locator))
            new_words += 1
            if len(window) == target_words:
                yield build_chunk()
                chunk_index += 1
                for _ in range(target_words - overlap_words):
                    window.popleft()
                new_words = 0
    if new_words:
        yield build_chunk()


def iter_document_units(path: str | Path, max_pages: int | None = None) -> Iterator[TextUnit]:
    path = Path(path)
    if path.stat().st_size > MAX_FILE_BYTES:
        raise PipelineError('FILE_TOO_LARGE', 'File exceeds the limit of 20 MiB.', 'extracting')
    content = path.read_bytes()
    doc = extract_document(content, path.suffix.lstrip('.'), max_pages=max_pages)
    for index, part in enumerate(doc.text.split('\n\n'), 1):
        yield TextUnit(part, page_number=index if doc.page_count else None, source_locator=f'part {index}')
