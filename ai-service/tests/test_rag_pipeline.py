import hashlib
import subprocess
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import rag_pipeline as pipeline
from rag_pipeline import (
    PipelineError, TextUnit, chunk_text, clean_text, extract_document,
    iter_chunks, iter_document_units,
)


def pdf_bytes(texts: list[str], encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({
            NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'),
            NameObject('/BaseFont'): NameObject('/Helvetica'),
        })
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): font}),
        })
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 20 20 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = stream
    if encrypted:
        writer.encrypt('test-only')
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_page_metadata_and_blank_pages(tmp_path):
    content = pdf_bytes(['first page', '', 'third page'])
    result = extract_document(content, '.PDF')
    assert result.page_count == 3
    assert '[Page 3]\nthird page' in result.text
    path = tmp_path / 'source.pdf'
    path.write_bytes(content)
    units = list(iter_document_units(path))
    assert [unit.page_number for unit in units] == [1, 2, 3]
    chunks = list(iter_chunks(units))
    assert chunks[0]['page_numbers'] == [1, 3]
    assert extract_document(content, 'pdf', max_pages=1).page_count == 1


@pytest.mark.parametrize(('content', 'extension', 'code'), [
    (b'broken', 'pdf', 'PDF_EXTRACTION_FAILED'),
    (b'broken', 'docx', 'DOCX_EXTRACTION_FAILED'),
    (b'broken', 'exe', 'UNSUPPORTED_FILE_TYPE'),
    (pdf_bytes(['']), 'pdf', 'NO_EXTRACTABLE_TEXT'),
    (pdf_bytes(['secret'], True), 'pdf', 'PDF_ENCRYPTED'),
])
def test_extraction_errors(content, extension, code):
    with pytest.raises(PipelineError) as caught:
        extract_document(content, extension)
    assert caught.value.code == code
    assert caught.value.stage == 'extracting'


def test_docx_keeps_document_order_and_does_not_invent_pages():
    document = Document()
    document.add_paragraph('Tiếng Việt lecturer@example.com')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Bảng ở giữa'
    document.add_paragraph('Đoạn cuối')
    output = BytesIO()
    document.save(output)
    result = extract_document(output.getvalue(), 'docx')
    assert result.text.index('Tiếng Việt') < result.text.index('Bảng ở giữa') < result.text.index('Đoạn cuối')
    assert 'lecturer@example.com' not in result.text
    assert result.page_count is None
    assert chunk_text(result.text)[0]['page_number'] is None


def test_empty_docx():
    output = BytesIO()
    Document().save(output)
    with pytest.raises(PipelineError, match='No text'):
        extract_document(output.getvalue(), 'docx')


def test_docx_zip_expansion_limit(monkeypatch):
    output = BytesIO()
    with ZipFile(output, 'w') as archive:
        archive.writestr('oversized.xml', 'x' * 20)
    monkeypatch.setattr(pipeline, 'MAX_DOCX_UNCOMPRESSED_BYTES', 10)
    with pytest.raises(PipelineError) as caught:
        extract_document(output.getvalue(), 'docx')
    assert caught.value.code == 'DOCX_TOO_LARGE'


def test_clean_preserves_vietnamese_math_and_redacts_emails():
    raw = '  Tiếng\tViệt\x00\u200b\ufffd\r\n\n\nA ∪ B ≤ 3  a.b+tag@Example.edu.vn; tên@trường.vn '
    assert clean_text(raw) == 'Tiếng Việt\n\nA ∪ B ≤ 3 [redacted-email]; [redacted-email]'
    assert clean_text(clean_text(raw)) == clean_text(raw)


@pytest.mark.parametrize('length', [0, 1, 49, 50, 499, 500, 501, 950, 1000, 1400])
def test_chunk_sizes_overlap_md5_and_no_duplicate_tail(length):
    words = [f'w{i}' for i in range(length)]
    chunks = chunk_text(' '.join(words))
    assert chunks == chunk_text(' '.join(words))
    restored = []
    for index, chunk in enumerate(chunks):
        tokens = chunk['content'].split()
        assert 0 < len(tokens) <= 500
        assert chunk['chunk_index'] == index
        assert chunk['content_hash'] == hashlib.md5(chunk['content'].encode()).hexdigest()
        assert chunk['hash_algorithm'] == 'md5'
        if index:
            assert tokens[:50] == chunks[index - 1]['content'].split()[-50:]
        restored.extend(tokens if index == 0 else tokens[50:])
    assert restored == words
    if length == 500:
        assert len(chunks) == 1


@pytest.mark.parametrize(('size', 'overlap'), [(0, 0), (-1, 0), (5, -1), (5, 5), (5, 6)])
def test_invalid_chunk_config(size, overlap):
    with pytest.raises(ValueError):
        chunk_text('hello', size, overlap)


def test_overlap_preserves_original_page_metadata():
    chunks = chunk_text('[Page 2]\na b c d\n[Page 5]\ne f g h', 5, 2)
    assert chunks[0]['page_numbers'] == [2, 5]
    assert chunks[1]['page_numbers'] == [2, 5]
    assert chunks[1]['page_number'] == 2
    assert chunks[1]['content'] == 'd e f g h'


def test_zero_overlap_and_stream_consumption():
    consumed = []
    def units():
        for index in range(100):
            consumed.append(index)
            yield TextUnit(f'w{index}', index + 1)
    chunks = iter_chunks(units(), 5, 0)
    assert next(chunks)['content'] == 'w0 w1 w2 w3 w4'
    assert consumed == list(range(5))
    assert next(chunks)['page_numbers'] == [6, 7, 8, 9, 10]


@pytest.mark.parametrize('mode', ['success', 'missing', 'timeout', 'bad', 'empty'])
def test_antiword_safe_process_and_cleanup(monkeypatch, mode):
    paths = []
    def run(args, **kwargs):
        paths.append(Path(args[-1]))
        assert paths[-1].read_bytes() == b'legacy'
        assert args[:3] == ['antiword', '-m', 'UTF-8.txt']
        assert kwargs['timeout'] == 30 and kwargs['shell'] is False
        if mode == 'missing':
            raise FileNotFoundError()
        if mode == 'timeout':
            raise subprocess.TimeoutExpired(args, 30)
        kwargs['stdout'].write('Tiếng Việt'.encode() if mode == 'success' else b'')
        return subprocess.CompletedProcess(args, 1 if mode == 'bad' else 0)
    monkeypatch.setattr(pipeline.subprocess, 'run', run)
    if mode == 'success':
        result = extract_document(b'legacy', 'doc')
        assert 'Tiếng Việt' in result.text
        assert result.page_count is None
    else:
        code = {'missing': 'DOC_PARSER_UNAVAILABLE', 'timeout': 'DOC_EXTRACTION_TIMEOUT',
                'bad': 'DOC_EXTRACTION_FAILED', 'empty': 'NO_EXTRACTABLE_TEXT'}[mode]
        with pytest.raises(PipelineError) as caught:
            extract_document(b'legacy', 'doc')
        assert caught.value.code == code
    assert paths and not paths[0].exists()


def test_limits_for_bytes_path_and_extracted_text(monkeypatch, tmp_path):
    monkeypatch.setattr(pipeline, 'MAX_FILE_BYTES', 3)
    with pytest.raises(PipelineError) as caught:
        extract_document(b'1234', 'pdf')
    assert caught.value.code == 'FILE_TOO_LARGE'
    path = tmp_path / 'large.pdf'
    path.write_bytes(b'1234')
    with pytest.raises(PipelineError):
        list(iter_document_units(path))
    monkeypatch.setattr(pipeline, 'MAX_FILE_BYTES', 100000)
    monkeypatch.setattr(pipeline, 'MAX_TEXT_CHARS', 3)
    with pytest.raises(PipelineError) as caught:
        extract_document(pdf_bytes(['too long']), 'pdf')
    assert caught.value.code == 'TEXT_TOO_LARGE'


def test_invalid_max_pages():
    with pytest.raises(ValueError):
        extract_document(pdf_bytes(['hello']), 'pdf', 0)
